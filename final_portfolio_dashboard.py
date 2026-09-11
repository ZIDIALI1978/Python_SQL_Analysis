import os
import shutil
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.utils import get_column_letter


# ============================================================
# 1. FILE PATHS
# ============================================================

base_folder = r"C:\Users\win\Desktop\sql\Python_SQL_Analysis"

base_file = os.path.join(
    base_folder,
    "Retail_ECommerce_Final_With_RFM.xlsx"
)

clv_file = os.path.join(
    base_folder,
    "CLV_Analysis.xlsx"
)

cohort_file = os.path.join(
    base_folder,
    "Cohort_Analysis.xlsx"
)

product_file = os.path.join(
    base_folder,
    "Product_Inventory_Analysis.xlsx"
)

supplier_file = os.path.join(
    base_folder,
    "Supplier_Operations_Analysis.xlsx"
)

output_file = os.path.join(
    base_folder,
    "Retail_ECommerce_Final_Portfolio.xlsx"
)


# ============================================================
# 2. CHECK FILES
# ============================================================

required_files = [
    base_file,
    clv_file,
    cohort_file,
    product_file,
    supplier_file
]

print("CHECKING REQUIRED FILES")

for file in required_files:

    if os.path.exists(file):

        print(
            "Found:",
            os.path.basename(file)
        )

    else:

        raise FileNotFoundError(
            f"Required file not found: {file}"
        )


# ============================================================
# 3. REMOVE OLD OUTPUT
# ============================================================

if os.path.exists(output_file):

    try:

        os.remove(output_file)

        print("\nOld output file removed.")

    except PermissionError:

        raise PermissionError(
            "Please close Retail_ECommerce_Final_Portfolio.xlsx "
            "in Excel and run the script again."
        )


# ============================================================
# 4. COPY BASE WORKBOOK
# ============================================================

shutil.copy2(
    base_file,
    output_file
)

print("Base workbook copied successfully.")


# ============================================================
# 5. OPEN WORKBOOK
# ============================================================

wb = load_workbook(output_file)

print(
    f"Workbook opened successfully. "
    f"Existing sheets: {len(wb.sheetnames)}"
)


# ============================================================
# 6. HELPER FUNCTION
# ============================================================

def import_excel_sheets(
    source_file,
    target_workbook,
    skip_sheets=None
):

    if skip_sheets is None:
        skip_sheets = []

    source_wb = load_workbook(
        source_file,
        data_only=False
    )

    for source_sheet_name in source_wb.sheetnames:

        if source_sheet_name in skip_sheets:
            continue

        source_ws = source_wb[
            source_sheet_name
        ]

        # Avoid duplicate sheet names
        target_name = source_sheet_name

        if target_name in target_workbook.sheetnames:

            counter = 2

            while (
                f"{source_sheet_name}_{counter}"
                in target_workbook.sheetnames
            ):

                counter += 1

            target_name = (
                f"{source_sheet_name}_{counter}"
            )

        target_ws = target_workbook.create_sheet(
            title=target_name
        )

        for row in source_ws.iter_rows():

            for cell in row:

                new_cell = target_ws.cell(
                    row=cell.row,
                    column=cell.column,
                    value=cell.value
                )

                if cell.number_format:
                    new_cell.number_format = (
                        cell.number_format
                    )

                if cell.alignment:
                    new_cell.alignment = (
                        cell.alignment.copy()
                    )

                if cell.font:
                    new_cell.font = (
                        cell.font.copy()
                    )

                if cell.fill:
                    new_cell.fill = (
                        cell.fill.copy()
                    )

        # Copy widths
        for column_letter, dimension in (
            source_ws.column_dimensions.items()
        ):

            target_ws.column_dimensions[
                column_letter
            ].width = dimension.width

        # Freeze panes
        target_ws.freeze_panes = (
            source_ws.freeze_panes
        )

        print(
            f"Imported sheet: {target_name}"
        )


# ============================================================
# 7. IMPORT CLV
# ============================================================

print("\nIMPORTING CLV ANALYSIS")

import_excel_sheets(
    clv_file,
    wb
)


# ============================================================
# 8. IMPORT COHORT
# ============================================================

print("\nIMPORTING COHORT ANALYSIS")

import_excel_sheets(
    cohort_file,
    wb
)


# ============================================================
# 9. IMPORT PRODUCT & INVENTORY
# ============================================================

print(
    "\nIMPORTING PRODUCT & INVENTORY ANALYSIS"
)

import_excel_sheets(
    product_file,
    wb
)


# ============================================================
# 10. IMPORT SUPPLIER & OPERATIONS
# ============================================================

print(
    "\nIMPORTING SUPPLIER & OPERATIONS ANALYSIS"
)

import_excel_sheets(
    supplier_file,
    wb
)


# ============================================================
# 11. CREATE EXECUTIVE DASHBOARD
# ============================================================

if "Executive Dashboard" in wb.sheetnames:

    del wb["Executive Dashboard"]

dashboard = wb.create_sheet(
    "Executive Dashboard",
    0
)

print("\nExecutive Dashboard created.")


# ============================================================
# 12. DASHBOARD STYLING
# ============================================================

title_fill = PatternFill(
    fill_type="solid",
    fgColor="1F4E78"
)

section_fill = PatternFill(
    fill_type="solid",
    fgColor="D9EAF7"
)

kpi_fill = PatternFill(
    fill_type="solid",
    fgColor="EAF2F8"
)

white_font = Font(
    color="FFFFFF",
    bold=True,
    size=16
)

section_font = Font(
    bold=True,
    size=12
)

kpi_title_font = Font(
    bold=True,
    size=10
)

kpi_value_font = Font(
    bold=True,
    size=16
)

thin_border = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin")
)


# ============================================================
# 13. DASHBOARD TITLE
# ============================================================

dashboard.merge_cells(
    "A1:L2"
)

dashboard["A1"] = (
    "RETAIL E-COMMERCE ANALYTICS"
)

dashboard["A1"].fill = title_fill
dashboard["A1"].font = white_font
dashboard["A1"].alignment = Alignment(
    horizontal="center",
    vertical="center"
)

dashboard.row_dimensions[1].height = 28
dashboard.row_dimensions[2].height = 28


dashboard.merge_cells(
    "A3:L3"
)

dashboard["A3"] = (
    "Executive Business Intelligence Dashboard"
)

dashboard["A3"].font = Font(
    bold=True,
    italic=True,
    size=11
)

dashboard["A3"].alignment = Alignment(
    horizontal="center"
)


# ============================================================
# 14. LOAD DATA FOR KPIs
# ============================================================

sales_master = pd.read_excel(
    base_file,
    sheet_name="Sales Master"
)

orders_data = pd.read_excel(
    base_file,
    sheet_name="Orders"
)

customers_data = pd.read_excel(
    base_file,
    sheet_name="Customers"
)

products_data = pd.read_excel(
    base_file,
    sheet_name="Products"
)

profit_data = pd.read_excel(
    base_file,
    sheet_name="Profit Analysis"
)


# ============================================================
# 15. KPI CALCULATIONS
# ============================================================

total_revenue = (
    sales_master["Revenue"].sum()
    if "Revenue" in sales_master.columns
    else 0
)

total_orders = (
    orders_data["OrderID"].nunique()
    if "OrderID" in orders_data.columns
    else 0
)

total_customers = (
    customers_data["CustomerID"].nunique()
    if "CustomerID" in customers_data.columns
    else 0
)

total_products = (
    products_data["ProductID"].nunique()
    if "ProductID" in products_data.columns
    else 0
)

total_quantity = (
    sales_master["Quantity"].sum()
    if "Quantity" in sales_master.columns
    else 0
)

total_profit = 0

if "Profit" in sales_master.columns:

    total_profit = sales_master["Profit"].sum()

elif "TotalProfit" in profit_data.columns:

    total_profit = profit_data["TotalProfit"].sum()


aov = (
    total_revenue / total_orders
    if total_orders > 0
    else 0
)

profit_margin = (
    total_profit / total_revenue * 100
    if total_revenue > 0
    else 0
)


# ============================================================
# 16. KPI CARDS
# ============================================================

kpis = [
    ("Total Revenue", total_revenue),
    ("Total Profit", total_profit),
    ("Total Orders", total_orders),
    ("Total Customers", total_customers),
    ("Total Products", total_products),
    ("Average Order Value", aov),
    ("Quantity Sold", total_quantity),
    ("Profit Margin", profit_margin)
]


kpi_positions = [
    ("A5:C7"),
    ("D5:F7"),
    ("G5:I7"),
    ("J5:L7"),
    ("A9:C11"),
    ("D9:F11"),
    ("G9:I11"),
    ("J9:L11")
]


for (
    (title, value),
    cell_range
) in zip(
    kpis,
    kpi_positions
):

    dashboard.merge_cells(
        cell_range
    )

    start_cell = cell_range.split(":")[0]

    cell = dashboard[start_cell]

    cell.value = (
        f"{title}\n\n"
        f"{value:,.2f}"
    )

    cell.fill = kpi_fill
    cell.border = thin_border
    cell.alignment = Alignment(
        horizontal="center",
        vertical="center",
        wrap_text=True
    )

    cell.font = kpi_value_font


# ============================================================
# 17. SECTION HEADERS
# ============================================================

sections = {
    "A13:F13": "SALES PERFORMANCE",
    "G13:L13": "CUSTOMER INTELLIGENCE",
    "A30:F30": "PRODUCT & INVENTORY",
    "G30:L30": "OPERATIONS"
}


for cell_range, title in sections.items():

    dashboard.merge_cells(
        cell_range
    )

    start_cell = cell_range.split(":")[0]

    cell = dashboard[start_cell]

    cell.value = title
    cell.fill = section_fill
    cell.font = section_font
    cell.alignment = Alignment(
        horizontal="center"
    )


# ============================================================
# 18. MONTHLY REVENUE DATA
# ============================================================

monthly_revenue = pd.read_excel(
    base_file,
    sheet_name="Monthly Revenue",
    header=None
)

print(
    "\nMonthly Revenue raw data:"
)

print(
    monthly_revenue.head(10)
)


# Find actual data rows
monthly_rows = []

for _, row in monthly_revenue.iterrows():

    values = row.tolist()

    if len(values) < 2:
        continue

    month = values[0]
    revenue = values[1]

    # Skip title/header rows
    if pd.isna(month) or pd.isna(revenue):
        continue

    # Skip text headers
    if str(month).upper() in [
        "MONTHLY REVENUE ANALYSIS",
        "MONTH",
        "MONTHLY REVENUE"
    ]:
        continue

    try:

        revenue = float(revenue)

        monthly_rows.append(
            (month, revenue)
        )

    except (
        ValueError,
        TypeError
    ):

        continue


print(
    f"Monthly revenue records found: "
    f"{len(monthly_rows)}"
)


# ============================================================
# 19. MONTHLY REVENUE TABLE
# ============================================================

monthly_start_row = 14
monthly_start_col = 1


# Remove any merged cells that overlap
# the monthly revenue table area

for merged_range in list(
    dashboard.merged_cells.ranges
):

    if (
        merged_range.min_row <=
        monthly_start_row + len(monthly_rows) + 2
        and
        merged_range.max_row >=
        monthly_start_row
        and
        merged_range.min_col <= 2
        and
        merged_range.max_col >= 1
    ):

        dashboard.unmerge_cells(
            str(merged_range)
        )


# Headers

dashboard.cell(
    monthly_start_row,
    monthly_start_col,
    "Month"
)

dashboard.cell(
    monthly_start_row,
    monthly_start_col + 1,
    "Revenue"
)


dashboard.cell(
    monthly_start_row,
    monthly_start_col
).font = Font(
    bold=True
)

dashboard.cell(
    monthly_start_row,
    monthly_start_col + 1
).font = Font(
    bold=True
)


# Write monthly data

for index, (
    month,
    revenue
) in enumerate(
    monthly_rows,
    start=monthly_start_row + 1
):

    dashboard.cell(
        index,
        monthly_start_col,
        month
    )

    dashboard.cell(
        index,
        monthly_start_col + 1,
        revenue
    )

    dashboard.cell(
        index,
        monthly_start_col + 1
    ).number_format = '#,##0.00'


# ============================================================
# 20. MONTHLY REVENUE CHART
# ============================================================

if len(monthly_rows) >= 2:

    line_chart = LineChart()

    line_chart.title = (
        "Monthly Revenue Trend"
    )

    line_chart.y_axis.title = (
        "Revenue"
    )

    line_chart.x_axis.title = (
        "Month"
    )

    data = Reference(
        dashboard,
        min_col=2,
        min_row=monthly_start_row,
        max_row=(
            monthly_start_row +
            len(monthly_rows)
        )
    )

    categories = Reference(
        dashboard,
        min_col=1,
        min_row=monthly_start_row + 1,
        max_row=(
            monthly_start_row +
            len(monthly_rows)
        )
    )

    line_chart.add_data(
        data,
        titles_from_data=True
    )

    line_chart.set_categories(
        categories
    )

    line_chart.height = 7
    line_chart.width = 12

    dashboard.add_chart(
        line_chart,
        "D14"
    )

    print(
        "Monthly Revenue chart added."
    )

else:

    print(
        "Monthly Revenue chart skipped: "
        "not enough data."
    )

# ============================================================
# 21. RFM SEGMENT DATA
# ============================================================

if "RFM Segment Counts" in wb.sheetnames:

    rfm_ws = wb[
        "RFM Segment Counts"
    ]

    dashboard["G14"] = (
        "RFM Segment"
    )

    dashboard["H14"] = (
        "Customers"
    )

    dashboard["G14"].font = Font(
        bold=True
    )

    dashboard["H14"].font = Font(
        bold=True
    )

    rfm_rows = []

    for row in rfm_ws.iter_rows(
        min_row=2,
        values_only=True
    ):

        if row[0] is not None:

            rfm_rows.append(
                (row[0], row[1])
            )

    for index, (
        segment,
        count
    ) in enumerate(
        rfm_rows,
        start=15
    ):

        dashboard.cell(
            index,
            7,
            segment
        )

        dashboard.cell(
            index,
            8,
            count
        )

    if rfm_rows:

        rfm_chart = BarChart()

        rfm_chart.title = (
            "Customer RFM Segments"
        )

        rfm_chart.y_axis.title = (
            "Customers"
        )

        rfm_data = Reference(
            dashboard,
            min_col=8,
            min_row=14,
            max_row=14 + len(rfm_rows)
        )

        rfm_categories = Reference(
            dashboard,
            min_col=7,
            min_row=15,
            max_row=14 + len(rfm_rows)
        )

        rfm_chart.add_data(
            rfm_data,
            titles_from_data=True
        )

        rfm_chart.set_categories(
            rfm_categories
        )

        rfm_chart.height = 7
        rfm_chart.width = 12

        dashboard.add_chart(
            rfm_chart,
            "J14"
        )


# ============================================================
# 22. PRODUCT & INVENTORY SUMMARY
# ============================================================

if "ABC Analysis" in wb.sheetnames:

    abc_ws = wb[
        "ABC Analysis"
    ]

    dashboard["A31"] = (
        "ABC Class"
    )

    dashboard["B31"] = (
        "Products"
    )

    dashboard["C31"] = (
        "Revenue %"
    )

    dashboard["A31"].font = Font(
        bold=True
    )

    dashboard["B31"].font = Font(
        bold=True
    )

    dashboard["C31"].font = Font(
        bold=True
    )

    row_number = 32

    for row in abc_ws.iter_rows(
        min_row=2,
        values_only=True
    ):

        if row[0] is None:
            continue

        dashboard.cell(
            row_number,
            1,
            row[0]
        )

        dashboard.cell(
            row_number,
            2,
            row[1]
        )

        dashboard.cell(
            row_number,
            3,
            row[5]
            if len(row) > 5
            else None
        )

        row_number += 1


# ============================================================
# 23. INVENTORY RISK
# ============================================================

if "Inventory Risk" in wb.sheetnames:

    inventory_ws = wb[
        "Inventory Risk"
    ]

    dashboard["A37"] = (
        "Inventory Risk"
    )

    dashboard["B37"] = (
        "Products"
    )

    dashboard["C37"] = (
        "Inventory Value"
    )

    dashboard["A37"].font = Font(
        bold=True
    )

    dashboard["B37"].font = Font(
        bold=True
    )

    dashboard["C37"].font = Font(
        bold=True
    )

    row_number = 38

    for row in inventory_ws.iter_rows(
        min_row=2,
        values_only=True
    ):

        if row[0] is None:
            continue

        dashboard.cell(
            row_number,
            1,
            row[0]
        )

        dashboard.cell(
            row_number,
            2,
            row[1]
        )

        dashboard.cell(
            row_number,
            3,
            row[2]
        )

        dashboard.cell(
            row_number,
            3
        ).number_format = '#,##0.00'

        row_number += 1


# ============================================================
# 24. OPERATIONS SUMMARY
# ============================================================

if "Delivery Summary" in wb.sheetnames:

    delivery_ws = wb[
        "Delivery Summary"
    ]

    dashboard["G31"] = (
        "Delivery KPI"
    )

    dashboard["H31"] = (
        "Value"
    )

    dashboard["G31"].font = Font(
        bold=True
    )

    dashboard["H31"].font = Font(
        bold=True
    )

    row_number = 32

    for row in delivery_ws.iter_rows(
        min_row=2,
        values_only=True
    ):

        if row[0] is None:
            continue

        dashboard.cell(
            row_number,
            7,
            row[0]
        )

        dashboard.cell(
            row_number,
            8,
            row[1]
        )

        row_number += 1


# ============================================================
# 25. COURIER PERFORMANCE
# ============================================================

if "Courier Performance" in wb.sheetnames:

    courier_ws = wb[
        "Courier Performance"
    ]

    dashboard["J31"] = (
        "Courier"
    )

    dashboard["K31"] = (
        "Shipments"
    )

    dashboard["L31"] = (
        "Avg Delivery Days"
    )

    dashboard["J31"].font = Font(
        bold=True
    )

    dashboard["K31"].font = Font(
        bold=True
    )

    dashboard["L31"].font = Font(
        bold=True
    )

    row_number = 32

    for row in courier_ws.iter_rows(
        min_row=2,
        values_only=True
    ):

        if row[0] is None:
            continue

        dashboard.cell(
            row_number,
            10,
            row[0]
        )

        dashboard.cell(
            row_number,
            11,
            row[1]
        )

        dashboard.cell(
            row_number,
            12,
            row[2]
        )

        dashboard.cell(
            row_number,
            12
        ).number_format = '0.00'

        row_number += 1


# ============================================================
# 26. BUSINESS INSIGHTS
# ============================================================

dashboard.merge_cells(
    "A45:L45"
)

dashboard["A45"] = (
    "KEY BUSINESS INSIGHTS"
)

dashboard["A45"].fill = section_fill
dashboard["A45"].font = section_font
dashboard["A45"].alignment = Alignment(
    horizontal="center"
)


insights = [
    (
        "1.",
        "Revenue concentration",
        "ABC-A products generate approximately "
        "80% of total revenue."
    ),
    (
        "2.",
        "Inventory risk",
        "Most products are classified as "
        "overstock based on current stock coverage."
    ),
    (
        "3.",
        "Delivery performance",
        "402 shipments have delivery dates, "
        "while 25 shipments remain pending."
    ),
    (
        "4.",
        "Average delivery time",
        "The observed average delivery time "
        "is approximately 4.1 days."
    ),
    (
        "5.",
        "Customer intelligence",
        "RFM and CLV analysis identify customers "
        "with different value and retention needs."
    )
]


row_number = 47

for number, title, description in insights:

    dashboard.cell(
        row_number,
        1,
        f"{number} {title}"
    )

    dashboard.merge_cells(
        start_row=row_number,
        start_column=2,
        end_row=row_number,
        end_column=12
    )

    dashboard.cell(
        row_number,
        2,
        description
    )

    dashboard.cell(
        row_number,
        1
    ).font = Font(
        bold=True
    )

    dashboard.cell(
        row_number,
        2
    ).alignment = Alignment(
        wrap_text=True
    )

    row_number += 2


# ============================================================
# 27. DASHBOARD WIDTHS
# ============================================================

widths = {
    "A": 22,
    "B": 18,
    "C": 18,
    "D": 18,
    "E": 18,
    "F": 18,
    "G": 22,
    "H": 18,
    "I": 18,
    "J": 22,
    "K": 18,
    "L": 20
}


for column, width in widths.items():

    dashboard.column_dimensions[
        column
    ].width = width


# ============================================================
# 28. FREEZE PANES
# ============================================================

dashboard.freeze_panes = "A5"


# ============================================================
# 29. ADD FOOTER
# ============================================================

dashboard.merge_cells(
    "A60:L60"
)

dashboard["A60"] = (
    "Retail E-Commerce Analytics | "
    "Python + SQL Server + Excel"
)

dashboard["A60"].alignment = Alignment(
    horizontal="center"
)

dashboard["A60"].font = Font(
    italic=True,
    size=10
)


# ============================================================
# 30. MOVE IMPORTANT SHEETS
# ============================================================

# Put Executive Dashboard first
wb._sheets.remove(
    dashboard
)

wb._sheets.insert(
    0,
    dashboard
)


# ============================================================
# 31. SAVE WORKBOOK
# ============================================================

wb.save(output_file)

print("\n" + "=" * 60)
print("FINAL PORTFOLIO DASHBOARD COMPLETED")
print("=" * 60)

print(
    f"Output file:\n{output_file}"
)

print(
    f"Total sheets: {len(wb.sheetnames)}"
)

print(
    "\nExecutive Dashboard added successfully."
)

print(
    "CLV Analysis imported successfully."
)

print(
    "Cohort Analysis imported successfully."
)

print(
    "Product & Inventory Analysis imported successfully."
)

print(
    "Supplier & Operations Analysis imported successfully."
)