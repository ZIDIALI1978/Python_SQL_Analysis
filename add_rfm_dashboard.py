import openpyxl
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
import os
import shutil


# ============================================================
# STEP 1: FILE PATHS
# ============================================================

input_file = "Retail_ECommerce_Final_With_RFM.xlsx"

output_file = "Retail_ECommerce_Final_With_RFM_Dashboard_v2.xlsx"


# ============================================================
# STEP 2: CHECK INPUT FILE
# ============================================================

if not os.path.exists(input_file):

    raise FileNotFoundError(
        f"Input file not found: {input_file}"
    )

print("Input Excel file found successfully.")


# ============================================================
# STEP 3: REMOVE OLD OUTPUT IF EXISTS
# ============================================================

if os.path.exists(output_file):

    try:

        os.remove(output_file)

        print("Old v2 output file removed successfully.")

    except PermissionError:

        raise PermissionError(
            "Please close Retail_ECommerce_Final_With_RFM_Dashboard_v2.xlsx "
            "from Excel and run again."
        )


# ============================================================
# STEP 4: COPY ORIGINAL RFM WORKBOOK
# ============================================================

try:

    shutil.copy2(
        input_file,
        output_file
    )

    print("Original RFM workbook copied successfully.")

except PermissionError:

    raise PermissionError(
        "The input Excel file is currently open. "
        "Please close Excel and run again."
    )


# ============================================================
# STEP 5: OPEN WORKBOOK
# ============================================================

wb = openpyxl.load_workbook(
    output_file
)

print("Workbook opened successfully.")

print(
    "Total sheets:",
    len(wb.sheetnames)
)


# ============================================================
# STEP 6: CHECK REQUIRED SHEETS
# ============================================================

if "Dashboard" not in wb.sheetnames:

    raise ValueError(
        "Dashboard sheet not found."
    )


if "RFM Segment Counts" not in wb.sheetnames:

    raise ValueError(
        "RFM Segment Counts sheet not found. "
        "Please run rfm_analysis.py first."
    )


dashboard = wb["Dashboard"]

rfm_sheet = wb["RFM Segment Counts"]


print("Dashboard sheet found.")

print("RFM Segment Counts sheet found.")


# ============================================================
# STEP 7: READ RFM DATA
# ============================================================

rfm_data = []

for row in rfm_sheet.iter_rows(
    min_row=2,
    values_only=True
):

    segment = row[0]
    count = row[1]

    if segment is not None and count is not None:

        rfm_data.append(
            (
                segment,
                count
            )
        )


print("RFM data loaded:")

print(rfm_data)


# ============================================================
# STEP 8: RFM SECTION LOCATION
# ============================================================

# IMPORTANT:
# Existing Dashboard charts are floating objects.
# Therefore max_row does not tell us where charts end.
#
# We intentionally place the RFM section much lower
# so it does not overlap existing charts.

start_row = 55

print(
    "RFM section will start at row:",
    start_row
)


# ============================================================
# STEP 9: RFM TITLE
# ============================================================

title_row = start_row

dashboard.merge_cells(
    start_row=title_row,
    start_column=1,
    end_row=title_row,
    end_column=8
)


title_cell = dashboard.cell(
    row=title_row,
    column=1
)


title_cell.value = (
    "CUSTOMER RFM SEGMENTATION"
)


title_cell.font = Font(
    bold=True,
    size=18
)


title_cell.alignment = Alignment(
    horizontal="center",
    vertical="center"
)


dashboard.row_dimensions[
    title_row
].height = 30


# ============================================================
# STEP 10: DESCRIPTION
# ============================================================

description_row = title_row + 1


dashboard.merge_cells(
    start_row=description_row,
    start_column=1,
    end_row=description_row,
    end_column=8
)


description_cell = dashboard.cell(
    row=description_row,
    column=1
)


description_cell.value = (
    "Customer segmentation based on Recency, "
    "Frequency and Monetary value."
)


description_cell.font = Font(
    italic=True,
    size=10
)


description_cell.alignment = Alignment(
    horizontal="center"
)


# ============================================================
# STEP 11: TABLE HEADER
# ============================================================

table_header_row = title_row + 3


dashboard.cell(
    row=table_header_row,
    column=1
).value = "Segment"


dashboard.cell(
    row=table_header_row,
    column=2
).value = "Customers"


header_fill = PatternFill(
    fill_type="solid",
    fgColor="1F4E78"
)


header_font = Font(
    bold=True,
    color="FFFFFF"
)


for column in range(1, 3):

    cell = dashboard.cell(
        row=table_header_row,
        column=column
    )

    cell.fill = header_fill

    cell.font = header_font

    cell.alignment = Alignment(
        horizontal="center",
        vertical="center"
    )


# ============================================================
# STEP 12: WRITE RFM TABLE
# ============================================================

data_start_row = table_header_row + 1


for index, (
    segment,
    count
) in enumerate(
    rfm_data,
    start=data_start_row
):

    dashboard.cell(
        row=index,
        column=1
    ).value = segment


    dashboard.cell(
        row=index,
        column=2
    ).value = count


# ============================================================
# STEP 13: TABLE FORMATTING
# ============================================================

thin_border = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin")
)


last_data_row = (
    data_start_row +
    len(rfm_data) -
    1
)


for row in dashboard.iter_rows(
    min_row=table_header_row,
    max_row=last_data_row,
    min_col=1,
    max_col=2
):

    for cell in row:

        cell.border = thin_border

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )


# ============================================================
# STEP 14: COLUMN WIDTH
# ============================================================

dashboard.column_dimensions[
    "A"
].width = 24


dashboard.column_dimensions[
    "B"
].width = 15


# ============================================================
# STEP 15: CREATE RFM BAR CHART
# ============================================================

chart = BarChart()


chart.type = "col"


chart.title = (
    "Customers by RFM Segment"
)


chart.y_axis.title = (
    "Number of Customers"
)


chart.x_axis.title = (
    "Customer Segment"
)


chart.height = 8


chart.width = 14


# ============================================================
# STEP 16: CHART DATA
# ============================================================

data = Reference(
    dashboard,
    min_col=2,
    min_row=table_header_row,
    max_row=last_data_row
)


categories = Reference(
    dashboard,
    min_col=1,
    min_row=data_start_row,
    max_row=last_data_row
)


chart.add_data(
    data,
    titles_from_data=True
)


chart.set_categories(
    categories
)


# ============================================================
# STEP 17: PLACE CHART
# ============================================================

chart_position = (
    "D" +
    str(table_header_row)
)


dashboard.add_chart(
    chart,
    chart_position
)


print(
    "RFM chart added at:",
    chart_position
)


# ============================================================
# STEP 18: RFM BUSINESS INSIGHTS
# ============================================================

insight_title_row = (
    last_data_row +
    3
)


dashboard.merge_cells(
    start_row=insight_title_row,
    start_column=1,
    end_row=insight_title_row,
    end_column=8
)


insight_title = dashboard.cell(
    row=insight_title_row,
    column=1
)


insight_title.value = (
    "RFM BUSINESS INSIGHTS"
)


insight_title.font = Font(
    bold=True,
    size=15
)


insight_title.alignment = Alignment(
    horizontal="center"
)


# ============================================================
# STEP 19: CREATE SEGMENT DICTIONARY
# ============================================================

segment_dict = {
    segment: count
    for segment, count in rfm_data
}


champions = segment_dict.get(
    "Champions",
    0
)


loyal = segment_dict.get(
    "Loyal",
    0
)


new_customers = segment_dict.get(
    "New",
    0
)


at_risk = segment_dict.get(
    "At Risk",
    0
)


lost = segment_dict.get(
    "Lost",
    0
)


# ============================================================
# STEP 20: BUSINESS INSIGHT 1
# ============================================================

insight_row = (
    insight_title_row +
    2
)


dashboard.merge_cells(
    start_row=insight_row,
    start_column=1,
    end_row=insight_row,
    end_column=8
)


dashboard.cell(
    row=insight_row,
    column=1
).value = (
    f"• Champions: {champions} customers. "
    "These are high-value customers and should be retained."
)


# ============================================================
# STEP 21: BUSINESS INSIGHT 2
# ============================================================

insight_row += 1


dashboard.merge_cells(
    start_row=insight_row,
    start_column=1,
    end_row=insight_row,
    end_column=8
)


dashboard.cell(
    row=insight_row,
    column=1
).value = (
    f"• Loyal: {loyal} customers. "
    "Maintain engagement and encourage repeat purchases."
)


# ============================================================
# STEP 22: BUSINESS INSIGHT 3
# ============================================================

insight_row += 1


dashboard.merge_cells(
    start_row=insight_row,
    start_column=1,
    end_row=insight_row,
    end_column=8
)


dashboard.cell(
    row=insight_row,
    column=1
).value = (
    f"• New: {new_customers} customers. "
    "Focus on onboarding and second-purchase conversion."
)


# ============================================================
# STEP 23: BUSINESS INSIGHT 4
# ============================================================

insight_row += 1


dashboard.merge_cells(
    start_row=insight_row,
    start_column=1,
    end_row=insight_row,
    end_column=8
)


dashboard.cell(
    row=insight_row,
    column=1
).value = (
    f"• At Risk: {at_risk} customers. "
    "Retention campaigns should target this group."
)


# ============================================================
# STEP 24: BUSINESS INSIGHT 5
# ============================================================

insight_row += 1


dashboard.merge_cells(
    start_row=insight_row,
    start_column=1,
    end_row=insight_row,
    end_column=8
)


dashboard.cell(
    row=insight_row,
    column=1
).value = (
    f"• Lost: {lost} customers. "
    "Win-back campaigns can be used to reactivate them."
)


# ============================================================
# STEP 25: FORMAT INSIGHTS
# ============================================================

for row in range(
    insight_title_row + 2,
    insight_row + 1
):

    cell = dashboard.cell(
        row=row,
        column=1
    )


    cell.alignment = Alignment(
        wrap_text=True,
        vertical="center"
    )


    dashboard.row_dimensions[
        row
    ].height = 25


# ============================================================
# STEP 26: SAVE
# ============================================================

wb.save(
    output_file
)


# ============================================================
# STEP 27: FINAL MESSAGE
# ============================================================

print()
print("=" * 55)
print("RFM DASHBOARD UPDATE COMPLETED SUCCESSFULLY")
print("=" * 55)

print(
    "Output file:",
    output_file
)

print(
    "Total sheets:",
    len(wb.sheetnames)
)

print(
    "RFM section starts at row:",
    start_row
)

print(
    "RFM chart position:",
    chart_position
)

print(
    "Existing sheets preserved."
)

print(
    "No previous RFM dashboard section copied."
)

print("=" * 55)